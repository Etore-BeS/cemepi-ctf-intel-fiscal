"""Export outliers review queue CSV to a researcher-friendly XLSX workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import polars as pl
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config.paths import PIPELINE_ROOT

DEFAULT_CSV = PIPELINE_ROOT / "notebooks/playground/output/outliers_tempo_valor_queue.csv"
DEFAULT_STATS = (
    PIPELINE_ROOT / "notebooks/playground/output/outliers_tempo_valor_stats.json"
)
DEFAULT_XLSX = (
    PIPELINE_ROOT / "notebooks/playground/output/outliers_tempo_valor_revisao.xlsx"
)

REVIEW_COLUMNS = [
    "status_revisao",
    "modus_operandi",
    "postura_juizo",
    "padrao_credor_fesp",
    "padrao_tramitacao",
    "resultado_qualitativo",
    "notas_pesquisador",
]

MAIN_COLUMNS = [
    "fila_pos",
    "outlier_tipo",
    "numero",
    "valor_corrigido_atual",
    "tempo_ate_sentenca_meses",
    "score_valor_x_tempo",
    "foro",
    "vara",
    "autores",
    "reus",
    "assunto",
    "tipo_ato",
    "tipo_sentença",
    "magistrado",
    "distribuicao_data",
    "data_sentenca_clean",
    "data_disponibilizacao",
    "decisao_len",
    "decisao_preview",
    *REVIEW_COLUMNS,
    "cd_processo",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INSTR_FILL = PatternFill("solid", fgColor="E8F0FE")
AMBIOS_FILL = PatternFill("solid", fgColor="FFF3CD")


def _write_sheet_rows(
    ws, rows: list[list], *, col_widths: dict[int, float] | None = None
) -> None:
    for row in rows:
        ws.append(row)
    if col_widths:
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width


def _style_header_row(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _build_instructions_rows() -> list[list]:
    return [
        ["Revisão qualitativa — outliers tempo × valor (FESP × ICMS)"],
        [""],
        ["Objetivo"],
        [
            "Identificar modus operandi na atuação jurisdicional em execuções fiscais "
            "estaduais paulistas com valor e/ou duração atípicos."
        ],
        [""],
        ["Como usar este arquivo"],
        ["1.", "Leia a aba Legenda antes de codificar."],
        [
            "2.",
            "Trabalhe na aba Fila_revisao — comece pelos registros outlier_tipo = ambos.",
        ],
        ["3.", "Preencha as colunas de revisão (azul claro no cabeçalho)."],
        [
            "4.",
            "Use decisao_preview na fila; texto completo na aba Texto_integral (mesmo fila_pos).",
        ],
        ["5.", "Filtros e ordenação do Excel estão habilitados na linha de cabeçalho."],
        [
            "6.",
            "status_revisao: use o dropdown (pendente → em_analise → concluido ou dispensado).",
        ],
        [""],
        ["Fluxo sugerido por processo"],
        ["a)", "Ler decisao_preview / Texto_integral"],
        ["b)", "Anotar modus_operandi (códigos curtos separados por ;)"],
        ["c)", "Classificar postura_juizo, padrao_credor_fesp, padrao_tramitacao"],
        ["d)", "Registrar resultado_qualitativo e notas livres"],
        ["e)", "Marcar status_revisao = concluido"],
        [""],
        ["Priorização automática da fila"],
        [
            "",
            "1º ambos (valor E tempo) → 2º maior score_valor_x_tempo → demais outliers",
        ],
        [""],
        ["Limitações"],
        [
            "decisao vem do CJPG (publicações), não do PDF integral dos autos. "
            "Ausência de texto não invalida o outlier — pode não haver decisão indexada."
        ],
    ]


def _build_legend_rows(stats: dict) -> list[list]:
    return [
        ["Legenda de campos e códigos"],
        [""],
        ["Coluna", "Descrição"],
        ["fila_pos", "Posição na fila priorizada (0 = maior prioridade)"],
        ["outlier_tipo", "ambos | valor | tempo — ver limiares IQR abaixo"],
        ["numero", "Número CNJ do processo"],
        ["valor_corrigido_atual", "Valor da causa corrigido pelo IPCA (R$)"],
        ["tempo_ate_sentenca_meses", "Meses entre distribuição e sentença (FACE)"],
        [
            "score_valor_x_tempo",
            "Produto valor × tempo — desempate dentro do mesmo tipo",
        ],
        ["decisao_preview", "Primeiros ~400 caracteres da decisão mais longa no lake"],
        ["decisao_len", "Tamanho em caracteres da decisão escolhida (0 = sem texto)"],
        ["tipo_ato", "Extraído da primeira linha da decisão (ex.: SENTENÇA P)"],
        [""],
        ["Limiares IQR (amostra EPED FESP×ICMS)"],
        ["Métrica", "Valor"],
        ["Processos na amostra", stats.get("n_total", "—")],
        ["Outliers valor (>", f"R$ {stats.get('limite_valor_iqr', 0):,.2f}"],
        ["Outliers tempo (>", f"{stats.get('limite_tempo_iqr_meses', 0):.1f} meses"],
        ["Outliers ambos", stats.get("n_outlier_ambos", "—")],
        ["Fila exportada", stats.get("queue_exportada", "—")],
        [""],
        ["Colunas de revisão (preencher manualmente)"],
        ["status_revisao", "pendente | em_analise | concluido | dispensado"],
        [
            "modus_operandi",
            "Códigos curtos do padrão observado (ex.: extincao_rapida;penhora_repetida)",
        ],
        [
            "postura_juizo",
            "Ex.: rigoroso_default | abertura_embargos | extincao_sem_merito",
        ],
        ["padrao_credor_fesp", "Ex.: peticao_padrao | pedido_excepcional | inercia"],
        ["padrao_tramitacao", "Ex.: citacao_frustrada | acordo_tardio | massa_falida"],
        [
            "resultado_qualitativo",
            "Ex.: extincao_pagamento | prescricao | nulidade_cda",
        ],
        ["notas_pesquisador", "Texto livre — citações, páginas, hipóteses"],
        [""],
        ["Sugestão de códigos modus_operandi (não exaustivo)"],
        ["Código", "Quando usar"],
        ["extincao_rapida", "Encerramento precoce sem litígio prolongado"],
        ["litigio_prolongado", "Tramitação com múltiplas fases e longa duração"],
        ["penhora_repetida", "Várias tentativas de constrição patrimonial"],
        ["acordo_tardio", "Transação/remissão após anos de tramitação"],
        ["massa_falida", "Contexto de falência/recuperação do executado"],
        ["valor_discrepante", "Valor da causa incompatível com o mérito narrado"],
        [
            "sem_texto_publicado",
            "Outlier sem decisão no lake — revisar fonte alternativa",
        ],
    ]


def export_xlsx(
    *,
    csv_path: Path = DEFAULT_CSV,
    stats_path: Path = DEFAULT_STATS,
    xlsx_path: Path = DEFAULT_XLSX,
) -> Path:
    stats = (
        json.loads(stats_path.read_text(encoding="utf-8"))
        if stats_path.exists()
        else {}
    )

    df = pl.read_csv(
        csv_path,
        infer_schema_length=10_000,
        schema_overrides={"num_processo_limpo": pl.Utf8},
    )

    for col in REVIEW_COLUMNS:
        if col not in df.columns:
            df = df.with_columns(pl.lit(None).cast(pl.Utf8).alias(col))

    if "status_revisao" in df.columns:
        df = df.with_columns(
            pl.when(pl.col("status_revisao").is_null())
            .then(pl.lit("pendente"))
            .otherwise(pl.col("status_revisao"))
            .alias("status_revisao")
        )

    main_df = df.select([c for c in MAIN_COLUMNS if c in df.columns])
    texto_df = df.select(["fila_pos", "numero", "cd_processo", "decisao"])

    wb = Workbook()
    ws_instr = wb.active
    ws_instr.title = "Instruções"
    _write_sheet_rows(ws_instr, _build_instructions_rows(), col_widths={1: 100})
    ws_instr["A1"].font = Font(bold=True, size=14)
    for row in range(3, ws_instr.max_row + 1):
        if ws_instr.cell(row=row, column=1).value in {
            "Objetivo",
            "Como usar este arquivo",
            "Fluxo sugerido por processo",
            "Priorização automática da fila",
            "Limitações",
        }:
            ws_instr.cell(row=row, column=1).font = Font(bold=True)
            ws_instr.cell(row=row, column=1).fill = INSTR_FILL

    ws_leg = wb.create_sheet("Legenda")
    _write_sheet_rows(
        ws_leg,
        _build_legend_rows(stats),
        col_widths={1: 28, 2: 72},
    )
    ws_leg["A1"].font = Font(bold=True, size=14)
    _style_header_row(ws_leg, row=3)
    _style_header_row(ws_leg, row=16)

    ws_fila = wb.create_sheet("Fila_revisao")
    ws_fila.append(main_df.columns)
    for row in main_df.iter_rows():
        ws_fila.append(list(row))
    _style_header_row(ws_fila)

    review_start = main_df.columns.index("status_revisao") + 1
    review_end = main_df.columns.index("notas_pesquisador") + 1
    for col_idx in range(review_start, review_end + 1):
        ws_fila.cell(row=1, column=col_idx).fill = PatternFill(
            "solid", fgColor="D9EAF7"
        )

    widths = {
        1: 14,
        2: 10,
        3: 24,
        4: 16,
        5: 12,
        6: 14,
        7: 28,
        8: 28,
        9: 32,
        10: 32,
        11: 36,
        12: 14,
        13: 28,
        14: 22,
        15: 14,
        16: 14,
        17: 14,
        18: 10,
        19: 48,
    }
    for col_idx, width in widths.items():
        if col_idx <= ws_fila.max_column:
            ws_fila.column_dimensions[get_column_letter(col_idx)].width = width

    for col_idx in range(review_start, review_end + 1):
        letter = get_column_letter(col_idx)
        ws_fila.column_dimensions[letter].width = 22

    ws_fila.freeze_panes = "A2"
    ws_fila.auto_filter.ref = ws_fila.dimensions

    tipo_col = main_df.columns.index("outlier_tipo") + 1
    tipo_letter = get_column_letter(tipo_col)
    last_row = ws_fila.max_row
    ws_fila.conditional_formatting.add(
        f"{tipo_letter}2:{tipo_letter}{last_row}",
        FormulaRule(formula=[f'${tipo_letter}2="ambos"'], fill=AMBIOS_FILL),
    )

    status_col = main_df.columns.index("status_revisao") + 1
    status_letter = get_column_letter(status_col)
    dv = DataValidation(
        type="list",
        formula1='"pendente,em_analise,concluido,dispensado"',
        allow_blank=True,
    )
    dv.error = "Use: pendente, em_analise, concluido ou dispensado"
    dv.errorTitle = "Status inválido"
    ws_fila.add_data_validation(dv)
    dv.add(f"{status_letter}2:{status_letter}{last_row}")

    for row_idx in range(2, last_row + 1):
        ws_fila.row_dimensions[row_idx].height = 30
        for col_idx in range(1, ws_fila.max_column + 1):
            ws_fila.cell(row=row_idx, column=col_idx).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    ws_texto = wb.create_sheet("Texto_integral")
    ws_texto.append(texto_df.columns)
    for row in texto_df.iter_rows():
        ws_texto.append(list(row))
    _style_header_row(ws_texto)
    ws_texto.freeze_panes = "A2"
    ws_texto.column_dimensions["A"].width = 10
    ws_texto.column_dimensions["B"].width = 24
    ws_texto.column_dimensions["C"].width = 18
    ws_texto.column_dimensions["D"].width = 100
    for row_idx in range(2, ws_texto.max_row + 1):
        ws_texto.row_dimensions[row_idx].height = 120
        ws_texto.cell(row=row_idx, column=4).alignment = Alignment(
            vertical="top", wrap_text=True
        )

    wb.save(xlsx_path)
    return xlsx_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    parser.add_argument("--stats", type=Path, default=DEFAULT_STATS)
    parser.add_argument("--output", type=Path, default=DEFAULT_XLSX)
    args = parser.parse_args()
    out = export_xlsx(csv_path=args.csv, stats_path=args.stats, xlsx_path=args.output)
    print(f"XLSX exportado: {out}")


if __name__ == "__main__":
    main()
